import os
import json
from pymavlink import mavutil
import pandas as pd
from collections import defaultdict
import requests
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import json
import array

# Custom serialization function
def serialize_obj(obj):
    if isinstance(obj, (list, dict, str, int, float, bool, type(None))):
        return obj
    elif isinstance(obj, array.array):
        return list(obj)
    elif isinstance(obj, bytes):
        return obj.decode('utf-8', 'ignore')
    raise TypeError(f'Object of type {obj.__class__.__name__} is not JSON serializable')

def col_num_to_letter(n):
    """Convert a column number to an Excel-style column letter (e.g., 1 => 'A', 27 => 'AA')."""
    letter = ''
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        letter = chr(65 + remainder) + letter
    return letter

# Modified fetch_weather_data function
def fetch_weather_data(start_timestamp, end_timestamp):
    url = f"http://weather.jaedynchilton.com/historic?start-timestamp={int(start_timestamp)}&end-timestamp={int(end_timestamp)}"
    response = requests.get(url)
    if response.status_code == 200:
        # with open('weather_data_debug.json', 'w') as f:
        #     json.dump(response.json(), f, indent=4)
        return response.json()['data'], True
    else:
        print(f"Failed to fetch weather data: {response.status_code}")
        return [], False


def match_and_append_weather_data(df, weather_data):
    if not weather_data:
        return df  # Return the original DataFrame if no weather data is available

    weather_df = pd.DataFrame(weather_data)
    
    # Create a new column in the DataFrame to store the closest weather data timestamp
    closest_ts = []
    
    for flight_ts in df['Unix_Epoch_Time']:
        closest_weather_ts = min(weather_df['ts'], key=lambda x: abs(x - flight_ts))
        closest_ts.append(closest_weather_ts)
        
    df['closest_ts'] = closest_ts
    
    # Merge the DataFrame based on the closest timestamp
    merged_df = pd.merge(df, weather_df, left_on='closest_ts', right_on='ts', how='left')
    
    # Rename 'ts' to 'weather_ts'
    merged_df.rename(columns={'ts': 'weather_ts'}, inplace=True)
    
    # Drop the 'closest_ts' column as it's no longer needed
    merged_df.drop('closest_ts', axis=1, inplace=True)
    
    # Keep only the specified columns for weather data
    weather_columns_to_keep = ['bar_absolute', 'wind_speed_avg', 'wind_speed_hi', 'wind_speed_hi_at', 'wind_speed_hi_dir', 
                               'wind_dir_of_prevail', 'dew_point_last', 'hum_last', 'temp_avg', 'uv_index', 'wet_bulb_last', 'weather_ts']
    
    all_columns_to_keep = [col for col in merged_df.columns if col in df.columns or col in weather_columns_to_keep]
    return merged_df[all_columns_to_keep]

def calculate_unix_epoch_time(row):
    gps_week = row['GWk']
    gps_milliseconds = row['GMS']
    return (gps_week * 604800) + (gps_milliseconds / 1000) + 315964800

def calculate_unix_epoch_time_from_timeus(time_us, initial_unix_epoch_time, initial_time_us):
    delta_time_s = (time_us - initial_time_us) / 1e6  # Converting microseconds to seconds
    return initial_unix_epoch_time + delta_time_s

def add_table_to_sheet(book, sheet_name, num_columns, num_rows):
    """Add table formatting to a given Excel sheet."""
    sheet = book[sheet_name]
    end_col_letter = col_num_to_letter(num_columns)
    table_range = f"A1:{end_col_letter}{num_rows + 1}"  # +1 for the header
    table = Table(displayName=f"{sheet_name}Table", ref=table_range)
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=True
    )
    table.tableStyleInfo = style
    sheet.add_table(table)

# First, let's define a function to map the MODE numbers to their readable names.
# The mapping is specific to ArduPilot on helicopters. This info was sourced from ChatGPT-4. YMMV.
def map_mode_number_to_name(mode_number):
    mode_map = {
        0: 'STABILIZE',
        1: 'ACRO',
        2: 'ALT_HOLD',
        3: 'AUTO',
        4: 'GUIDED',
        5: 'LOITER',
        6: 'RTL',
        7: 'CIRCLE',
        9: 'LAND',
        11: 'DRIFT',
        13: 'SPORT',
        14: 'FLIP',
        15: 'AUTOTUNE',
        16: 'POSHOLD',
        17: 'BRAKE',
        18: 'THROW',
        19: 'AVOID_ADSB',
        20: 'GUIDED_NOGPS',
        21: 'SMART_RTL',
        22: 'FLOWHOLD',
        23: 'FOLLOW',
        24: 'ZIGZAG',
        25: 'SYSTEMID',
        26: 'AUTOROTATE'
    }
    return mode_map.get(mode_number, f'Unknown-{mode_number}')

def convert_bin_to_excel(bin_file_path, excel_file_path):
    allowed_packet_types = ["IMU", "MAG", "RCIN", "RCOU", "GPS", "GPA", "BAT", "POWR", "MCU", "BARO", "RATE", "ATT", "VIBE", "HELI"]
    mlog = mavutil.mavlink_connection(bin_file_path)
    parsed_mavlink_data = []
    instance_count = defaultdict(set)  # Keep track of instances for each packet type
    pretty_mavlink_data = []  # Step 1: Initialize an empty list
    mode_data_frames = []   
    
    # MAVLink Parsing and Identification of Multiple Instances
    while True:
        msg = mlog.recv_msg()
        if msg is None:
            break
        msg_dict = msg.to_dict()
        packet_type = msg_dict.get('mavpackettype')
        
        instance = msg_dict.get('I', 0)
        if packet_type == 'VIBE':
            instance = msg_dict.get('IMU', 0)
        if packet_type in allowed_packet_types:
            instance_count[packet_type].add(instance)
            parsed_mavlink_data.append(msg_dict)
        # Capture MODE messages and map to readable names
        if packet_type == 'MODE':
            mode_name = map_mode_number_to_name(msg_dict.get('Mode', -1))
            msg_dict['ModeName'] = mode_name
            mode_df = pd.DataFrame([msg_dict]).drop(columns=['mavpackettype'], errors='ignore')
            mode_data_frames.append(mode_df)
        # Step 2: Append pretty-printed JSON string to the list
        pretty_mavlink_data.append(json.dumps(msg_dict, indent=4, default=serialize_obj))

    # Step 3: Write the pretty-printed data to a file for debugging
    with open('debug_mavlink_data.json', 'w') as f:
        for item in pretty_mavlink_data:
            f.write(f"{item}\n")

    # Create DataFrames for allowed data types
    data_frames = {}
    for packet in parsed_mavlink_data:
        packet_type = packet.get('mavpackettype')
        instance = packet.get('I', 0)
        if packet_type == 'VIBE':
            instance = packet.get('IMU', 0)
        unique_packet_type = f"{packet_type}_{instance}" if len(instance_count[packet_type]) > 1 else packet_type
        df = pd.DataFrame([packet]).drop(columns=['mavpackettype'], errors='ignore')
        if unique_packet_type not in data_frames:
            data_frames[unique_packet_type] = df
        else:
            data_frames[unique_packet_type] = pd.concat([data_frames[unique_packet_type], df], ignore_index=True)

    # Add closest GPS time to each DataFrame
    gps_df = data_frames.get('GPS', pd.DataFrame())
    if not gps_df.empty:
        gps_df['Unix_Epoch_Time'] = gps_df.apply(calculate_unix_epoch_time, axis=1)
        initial_unix_epoch_time = gps_df['Unix_Epoch_Time'].iloc[0]
        initial_time_us = gps_df['TimeUS'].iloc[0]
        for packet_type, df in data_frames.items():
            df['Unix_Epoch_Time'] = df['TimeUS'].apply(
                lambda x: calculate_unix_epoch_time_from_timeus(x, initial_unix_epoch_time, initial_time_us)
            )

    # Add Unix_Epoch_Time calculation for MODE messages
    if mode_data_frames:
        all_mode_df = pd.concat(mode_data_frames, ignore_index=True)
        
        # Create the Unix_Epoch_Time column
        unix_epoch_time = all_mode_df['TimeUS'].apply(
            lambda x: calculate_unix_epoch_time_from_timeus(x, initial_unix_epoch_time, initial_time_us)
        )
        
        # Insert Unix_Epoch_Time as the first column
        all_mode_df.insert(0, 'Unix_Epoch_Time', unix_epoch_time)

    # Fetch Weather Data
    min_time = gps_df['Unix_Epoch_Time'].min() if not gps_df.empty else None
    max_time = gps_df['Unix_Epoch_Time'].max() + 900 if not gps_df.empty else None
    weather_data, _ = fetch_weather_data(min_time, max_time) if min_time and max_time else ([], False)
    
    # Create the "ALL" table
    all_df = gps_df.copy()
    time_window = 0.2

    # Define the desired order of appending packet types to the "ALL" table
    ordered_packet_types = ["GPS", "RATE", "RCIN", "RCOU", "VIBE_0", "VIBE_1", "HELI",
                            "IMU_0", "IMU_1", "POWR", "MCU", "BAT", "MAG_0", "MAG_1", 
                            "BARO", "ATT", "GPA"]

    # Rename GPS columns to include "GPS_" prefix
    gps_columns_to_rename = {col: f"GPS_{col}" for col in all_df.columns if col != 'Unix_Epoch_Time'}
    all_df.rename(columns=gps_columns_to_rename, inplace=True)
    
    # Append data in the specified order
    for packet_type in ordered_packet_types:
        if packet_type == 'GPS':
            continue
        if packet_type not in data_frames:
            continue
        df = data_frames[packet_type]
        avg_df = pd.DataFrame()
        
        for gps_time in all_df['Unix_Epoch_Time']:
            min_time = gps_time - time_window
            max_time = gps_time + time_window
            close_rows = df[(df['Unix_Epoch_Time'] >= min_time) & (df['Unix_Epoch_Time'] <= max_time)]
            avg_row = close_rows.mean()
            avg_row['Unix_Epoch_Time'] = gps_time
            avg_df = avg_df._append(avg_row, ignore_index=True)

        # Rename columns to include source table
        renamed_columns = {col: f"{packet_type}_{col}" for col in df.columns if col != 'Unix_Epoch_Time'}
        avg_df.rename(columns=renamed_columns, inplace=True)
        
        all_df = pd.merge(all_df, avg_df, on='Unix_Epoch_Time', how='left')

    # Append weather data to the "ALL" table
    if weather_data:
        all_df = match_and_append_weather_data(all_df, weather_data)

    # Reorder columns to ensure 'Unix_Epoch_Time' is the first column
    reordered_columns = ['Unix_Epoch_Time'] + [col for col in all_df.columns if col != 'Unix_Epoch_Time']
    all_df = all_df[reordered_columns]

    # Save as Excel
    with pd.ExcelWriter(excel_file_path, engine='openpyxl') as writer:
        all_df.to_excel(writer, sheet_name='ALL', index=False)
        for packet_type, df in data_frames.items():
            reordered_columns = ['Unix_Epoch_Time'] + [col for col in df.columns if col != 'Unix_Epoch_Time']
            df[reordered_columns].to_excel(writer, sheet_name=packet_type, index=False)
        all_mode_df.to_excel(writer, sheet_name='MODE', index=False)

    # Load the workbook and select the 'ALL' sheet
    book = load_workbook(excel_file_path)
    sheet = book['ALL']

    # Debugging: Print end_col_letter and len(all_df)
    end_col_letter = col_num_to_letter(len(all_df.columns))
    print("End Column Letter:", end_col_letter)
    print("Row Count:", len(all_df))

    # Create a table
    table_range = f"A1:{end_col_letter}{len(all_df) + 1}"  # +1 due to the header row
    print("Table Range:", table_range)

    table = Table(displayName="DataTable", ref=table_range)

    # Add a default style to the table
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    table.tableStyleInfo = style

    # Add table to the sheet
    sheet.add_table(table)

    # Save the changes
    book.save(excel_file_path)

# Your allowed_packet_types should be defined here.
# convert_bin_to_excel('2023-10-27 14-19-05.bin', 'your_output15.xlsx')
